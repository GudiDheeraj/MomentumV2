"""Lightweight plain-text → PDF / DOCX exporters.

Used for on-demand downloads of cover letters and (as a fallback) résumé
text in whichever format the user asks for. Kept deliberately simple: these
produce a clean, readable document from plain text rather than trying to
reproduce a designed layout. Both reportlab (PDF) and python-docx (DOCX)
are already project dependencies.
"""
import io
from typing import List, Optional
from xml.sax.saxutils import escape as _xml_escape

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph


def _blocks(text: str) -> List[str]:
    """Splits text into paragraph blocks on blank lines, preserving single
    newlines inside a block so short line breaks (e.g. an address) survive."""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []
    return [b.strip() for b in text.split("\n\n") if b.strip()]


def text_to_pdf_bytes(text: str, title: Optional[str] = None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=LETTER,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch,
        title=title or "Document",
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "DocBody", parent=styles["Normal"],
        fontName="Helvetica", fontSize=10.5, leading=15,
        alignment=TA_LEFT, spaceAfter=10,
    )
    head = ParagraphStyle(
        "DocHead", parent=styles["Title"],
        fontName="Helvetica-Bold", fontSize=15, spaceAfter=14, alignment=TA_LEFT,
    )
    flow = []
    if title:
        flow.append(Paragraph(_xml_escape(title), head))
    for block in _blocks(text):
        flow.append(Paragraph(_xml_escape(block).replace("\n", "<br/>"), body))
    if not flow:
        flow.append(Paragraph("(empty)", body))
    doc.build(flow)
    return buf.getvalue()


def text_to_docx_bytes(text: str, title: Optional[str] = None) -> bytes:
    import docx

    document = docx.Document()
    if title:
        document.add_heading(title, level=1)
    for block in _blocks(text):
        para = document.add_paragraph()
        for i, line in enumerate(block.split("\n")):
            if i:
                para.add_run().add_break()
            para.add_run(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
